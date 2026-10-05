"""Genera il modello del Google Sheet "NOVA - Database" (da importare in Google Drive).

Uso: python3 nova/src/foglio.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "foglio" / "NOVA - Database.xlsx"

SCHEDE = {
    "CASELLE": {
        "colonne": ["slot", "email", "ruolo", "attiva", "nome_mittente", "limite_cold_giorno",
                    "inizio_warmup", "giorno_warmup", "stato_warmup", "inviate_oggi", "data_conteggio", "note"],
        "righe": [
            [1, "tuonome.business@gmail.com", "INVIO", "SI", "Mario Rossi", 20, "", "", "", "", "",
             "Casella che spedisce le cold email. Collegala nel nodo 'Casella 1' del workflow 00."],
            [2, "tuonome.warmup@gmail.com", "WARMUP", "SI", "Luca Bianchi", "", "", "", "", "", "",
             "Casella che aiuta solo il warm-up. Collegala nel nodo 'Casella 2' del workflow 00."],
        ],
    },
    "CODA": {
        "colonne": ["lead_id", "stato", "nome", "categoria", "citta", "indirizzo", "telefono", "sito", "email",
                    "rating", "recensioni", "data_inserimento", "regola_id", "prova", "motivo", "casella",
                    "oggetto", "thread_id", "message_id", "data_invio", "data_followup1", "data_followup2",
                    "data_risposta", "data_demo", "note"],
        "righe": [],
    },
    "REGOLE": {
        "colonne": ["id", "attiva", "priorita", "problema", "cosa_cercare", "prova_richiesta", "frase_email", "note"],
        "righe": [
            ["R1", "SI", 1, "Il sito non si adatta allo smartphone",
             "Nei SEGNALI TECNICI la riga meta_viewport_mobile vale 'no'.",
             "Cita la riga 'meta_viewport_mobile: no'.",
             "Ho aperto {sito} dal telefono e ho notato che {prova} Oggi la maggior parte dei clienti cerca dallo smartphone, "
             "e un sito che non si adatta allo schermo rischia di far perdere contatti.",
             "ESEMPIO — sostituisci con le tue regole"],
            ["R2", "SI", 2, "Il sito sembra non aggiornato da anni",
             "Nei SEGNALI TECNICI anni_copyright contiene solo anni di almeno 3 anni fa (o nel testo ci sono date/eventi vecchi).",
             "Cita la riga anni_copyright oppure la frase del testo con la data vecchia.",
             "Visitando {sito} ho notato che {prova} Chi arriva sul sito potrebbe pensare che l'attività non sia più seguita.",
             "ESEMPIO — sostituisci con le tue regole"],
            ["R3", "SI", 3, "Manca un modo semplice per prenotare o contattare online",
             "prenotazione_online è 'no' e modulo_contatto è 'no' e whatsapp è 'no'.",
             "Cita le righe dei segnali tecnici che lo dimostrano (es. 'prenotazione_online: no').",
             "Sul sito di {nome_attivita} {prova} Molti clienti oggi vogliono prenotare o scrivere con un clic, "
             "anche fuori orario.",
             "ESEMPIO — sostituisci con le tue regole"],
            ["R4", "NO", 4, "Il sito non ha una descrizione per Google",
             "meta_description è '(assente)'.",
             "Cita la riga 'meta_description: (assente)'.",
             "Cercando {nome_attivita} su Google ho notato che {prova} È un dettaglio, ma incide su quante persone cliccano.",
             "ESEMPIO disattivato — metti SI per usarlo"],
        ],
    },
    "TESTI": {
        "colonne": ["chiave", "variante", "testo"],
        "righe": [
            ["oggetto", 1, "Una cosa che ho notato sul sito di {nome_attivita}"],
            ["oggetto", 2, "{nome_attivita}: una domanda sul sito"],
            ["oggetto", 3, "Domanda veloce su {sito}"],
            ["saluto", 1, "Buongiorno,"],
            ["saluto", 2, "Buongiorno, gentile team di {nome_attivita},"],
            ["apertura", 1, "mi chiamo {mittente} e lavoro con {azienda}: aiutiamo le attività di {citta} a ottenere più clienti dal proprio sito."],
            ["apertura", 2, "sono {mittente} di {azienda}. Le scrivo perché seguiamo diverse attività come la vostra a {citta}."],
            ["proposta", 1, "Se le fa piacere, posso prepararle gratuitamente una breve analisi con 2-3 interventi concreti. Le basta rispondere a questa mail."],
            ["proposta", 2, "Le andrebbe se le mando qualche idea concreta per sistemarlo? Nessun impegno: basta una risposta."],
            ["chiusura", 1, "Grazie per l'attenzione e buon lavoro."],
            ["chiusura", 2, "Resto a disposizione, buona giornata."],
            ["firma", 1, "{mittente}\n{azienda}\n{sito_azienda}"],
            ["opt_out", 1, "P.S. Ho trovato il vostro contatto pubblicato sul vostro sito. Se preferite non ricevere altre mail, "
                           "rispondete \"no grazie\" e non vi scriverò più."],
            ["followup_1", 1, "Buongiorno,\n\nle riscrivo brevemente nel caso la mia mail precedente si sia persa tra le altre. "
                              "L'analisi gratuita di {sito} è ancora disponibile: basta una risposta.\n\n{mittente}"],
            ["followup_2", 1, "Buongiorno,\n\nultimo messaggio da parte mia, promesso. Se in futuro vorrete migliorare il sito di "
                              "{nome_attivita}, mi trovate a questo indirizzo.\n\nBuon lavoro,\n{mittente}\n{azienda}"],
            ["demo_sollecito_1", 1, "Buongiorno,\n\nha avuto modo di dare un'occhiata alla demo? La trova qui: {link_demo}\n\n"
                                    "Se ha domande sono a disposizione.\n\n{mittente}"],
            ["demo_sollecito_2", 1, "Buongiorno,\n\nle lascio di nuovo il link della demo preparata per {nome_attivita}: {link_demo}\n\n"
                                    "Se vuole ne parliamo 10 minuti al telefono, quando le è più comodo.\n\n{mittente}\n{azienda}"],
        ],
    },
    "RISPOSTE": {
        "colonne": ["data", "tipo", "lead_id", "nome", "email_da", "casella", "oggetto", "estratto", "id_gmail", "thread_id"],
        "righe": [],
    },
    "REGISTRO": {
        "colonne": ["lead_id", "nome", "email", "casella", "thread_id", "data_demo", "link_demo", "stato",
                    "data_sollecito1", "data_sollecito2"],
        "righe": [],
    },
    "BLACKLIST": {
        "colonne": ["email", "dominio", "data", "motivo"],
        "righe": [],
    },
    "RICERCHE": {
        "colonne": ["data", "citta", "categoria"],
        "righe": [],
    },
}

LEGGIMI = [
    ["NOVA - Database", ""],
    ["", ""],
    ["Scheda", "A cosa serve"],
    ["CASELLE", "Le caselle Gmail (max 10). slot = numero del nodo 'Casella N' nel workflow 00. ruolo: INVIO (spedisce le cold email) "
                "o WARMUP (aiuta solo il warm-up). Le colonne da inizio_warmup in poi le compila il workflow 05."],
    ["CODA", "I lead trovati dal workflow 01. Non cambiare i nomi delle colonne."],
    ["REGOLE", "I problemi che sai risolvere. Claude ne cerca uno nel sito del lead e deve portare una citazione esatta. "
               "frase_email entra nella mail ({prova} = la frase scritta da Claude)."],
    ["TESTI", "Le frasi della mail. Più righe con la stessa chiave = varianti scelte a caso. "
              "Chiavi obbligatorie: oggetto, saluto, proposta, chiusura, firma, followup_1, followup_2 "
              "(+ demo_sollecito_1, demo_sollecito_2 se usi i workflow opzionali)."],
    ["RISPOSTE", "Risposte, bounce e richieste di non essere contattati (le scrive il workflow 03)."],
    ["REGISTRO", "Lead a cui hai mandato la demo (workflow opzionali 06 e 07)."],
    ["BLACKLIST", "Email o domini da non contattare mai. Puoi aggiungerne a mano."],
    ["RICERCHE", "Le ricerche su Google Maps già fatte (workflow 01)."],
    ["", ""],
    ["Segnaposto", "{nome_attivita} {categoria} {citta} {sito} {problema} {prova} {frase_problema} {mittente} {azienda} {sito_azienda} {link_demo}"],
    ["", ""],
    ["Stati della CODA", "NUOVO → INVIATA → FOLLOWUP 1 → FOLLOWUP 2 → NESSUNA RISPOSTA.  Altri: DA GUARDARE (prova non verificata), "
                         "RISPOSTO, BOUNCE, ESCLUSO, ERRORE INVIO, TEST (mail mandata in modalità test), DEMO INVIATA."],
    ["Rimettere in coda", "Per riprovare un lead in DA GUARDARE, TEST o ERRORE INVIO scrivi NUOVO nella colonna stato."],
]

INTESTAZIONE = PatternFill("solid", fgColor="1F3A5F")


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "LEGGIMI"
    for riga in LEGGIMI:
        ws.append(riga)
    ws["A1"].font = Font(bold=True, size=14)
    for c in ("A3", "B3"):
        ws[c].font = Font(bold=True)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 140
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for nome, dati in SCHEDE.items():
        ws = wb.create_sheet(nome)
        ws.append(dati["colonne"])
        for riga in dati["righe"]:
            ws.append(riga)
        for i, col in enumerate(dati["colonne"], start=1):
            cella = ws.cell(row=1, column=i)
            cella.font = Font(bold=True, color="FFFFFF")
            cella.fill = INTESTAZIONE
            larga = col in ("testo", "frase_email", "cosa_cercare", "prova_richiesta", "note", "prova", "motivo", "estratto", "problema")
            ws.column_dimensions[get_column_letter(i)].width = 60 if larga else 18
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"✔ {OUT.relative_to(OUT.parent.parent)}")


if __name__ == "__main__":
    main()
